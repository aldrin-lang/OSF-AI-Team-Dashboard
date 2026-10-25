"""Turn exported Google Sheets (.xlsx) into sheet_mirror JSON for the dashboard (one-time manual load).
usage: python3 scripts/xlsx-to-sheet-mirror.py <xlsx dir> <out dir>   (then: npx tsx scripts/load-sheet-mirror.mts <out dir>)
Each spreadsheet must be saved as <spreadsheetId>.xlsx. Values are shown the way the sheet formats them
(currency, %, dates), with colours, bold, merges, notes, frozen panes and hidden rows (marked hidden)."""
import sys, os, json, re, datetime, math
import openpyxl
from openpyxl.utils import get_column_letter

SRC_DIR, OUT = sys.argv[1], sys.argv[2]
os.makedirs(OUT, exist_ok=True)

MASTER = '1NpuXpWbA0OVUNFE5eKHZut0-6F4qE6kw6hODsmv9IDM'
ACC = '1ZO7ykIuKMzl2BLXo7kXbp5Q5cLfSmJJa_8JM6NvMTXU'
VAOPS = '1DUN82Fm6VbBOnDDglxLHv1ZeuaBDqx47_uY2KKosp2c'
FORECAST = '1gY9mE0DUl9fPYf4t4O4EopCfe1o4CwV8RSpzkCpyhpY'
MKT = '17ldN9KgiBZkRcyR5MPescWWODrHDkeide2BvOnp2eZg'
PROPFIX = '1XvGIPWgSSI2Nl8dJLSGoHk3ePzzI9LZVw3IiMVYxEL8'
ROEL = '1jxzMgBY61Q3fdTNZQ0-rwYPJIEKMYRRG'
NAMES = {MASTER: 'OSF_Mastersheet 2026', ACC: 'Customer Payments [New]', VAOPS: 'VA Ops File', FORECAST: 'VA Forecast',
         MKT: 'OSF Marketing Daily Numbers', PROPFIX: 'Propfix.ie Responses', ROEL: 'External Hires Update (Roel)'}

# [key, department, area, label, file id, tab, options] (same list as the Apps Script dashTabList_, minus the ads tracker)
TABS = [
  ['acc-client-payments', 'accounts', 'payments', 'Client VA payments', ACC, ' Client VA Payments'],
  ['acc-va-mom', 'accounts', 'payments', 'OSF VA MoM', ACC, 'OSF VA MOM v2'],
  ['acc-ai-mom', 'accounts', 'payments', 'OSF AI MoM', ACC, 'OSF AI Mom'],
  ['acc-summary', 'accounts', 'payments', 'Summary', ACC, 'SUMMARY'],
  ['acc-overdue', 'accounts', 'payments', 'Overdue list', ACC, 'OVERDUE LIST'],
  ['acc-osf-overdue', 'accounts', 'payments', 'OSF overdue', ACC, 'OSF Overdue list'],
  ['acc-active-paying', 'accounts', 'payments', 'Active paying', ACC, 'Active Paying'],
  ['acc-invoices', 'accounts', 'payments', 'Maria · invoices', ACC, 'MARIA - Invoices'],
  ['acc-auto-debit', 'accounts', 'payments', 'Auto debit', ACC, 'Auto Debit'],
  ['acc-hiring-fee', 'accounts', 'payments', 'VA hiring fee', ACC, 'VA HIRING FEE '],
  ['acc-pnl', 'accounts', 'payments', 'P&L', MASTER, '2. PNL'],
  ['acc-overheads', 'accounts', 'payments', 'Overheads', MASTER, '7. Overheads'],
  ['rec-screening-dashboard', 'recruitment', 'candidates', 'Screening dashboard', PROPFIX, 'Dashboard'],
  ['rec-new-applicants', 'recruitment', 'candidates', 'New applicants', PROPFIX, 'New Applicants/Screening', {'xlsx': 'New ApplicantsScreening'}],
  'ROEL',
  ['today-daily-report', 'today', 'reports', 'Morning report', MASTER, 'OSF Daily Report'],
  ['today-ai-tasks', 'today', 'reports', 'AI tasks', MASTER, 'AI Tasks'],
  ['today-daily-numbers', 'today', 'reports', 'Daily numbers', MASTER, '6. Daily Numbers Crunch'],
  ['today-executive', 'today', 'reports', 'Executive dashboard', MASTER, '1. Executive Dashboard'],
  ['today-weekly-meeting', 'today', 'reports', 'Weekly meeting', MASTER, 'Weekly Meeting'],
  ['int-interview-tracker', 'interviews', 'clients', 'Interview tracker', MASTER, 'Daily Interview Tracker'],
  ['int-booking-to-sales', 'interviews', 'clients', 'Booking to sales', FORECAST, 'Booking to Sales Form'],
  ['vam-call-logs', 'va-managers', 'clients', 'Client call logs', VAOPS, 'Client Call Logs'],
  ['vam-redflag', 'va-managers', 'clients', 'Red flags', VAOPS, 'Redflag'],
  ['vam-masterlist', 'va-managers', 'clients', 'VA masterlist', FORECAST, 'VA MASTERLIST'],
  ['vam-for-ops', 'va-managers', 'clients', 'For ops team', FORECAST, 'FOR OPS TEAM'],
  ['vam-manager-breakdown', 'va-managers', 'clients', 'Manager breakdown', FORECAST, 'Manager Breakdown'],
  ['vam-retention', 'va-managers', 'clients', 'Retention plan', MASTER, '3. Retention Plan Action List'],
  ['vam-ops-kpi', 'va-managers', 'clients', 'Ops KPI', MASTER, 'OPs KPI'],
  ['vam-concerns', 'va-managers', 'clients', 'Concerns 2.0', VAOPS, 'Concerns 2.0'],
  ['vam-replacements', 'va-managers', 'clients', 'Replacement / cancellation', VAOPS, 'Replacement/Cancellation', {'xlsx': 'ReplacementCancellation'}],
  ['vam-cancelled', 'va-managers', 'clients', 'Cancelled campaigns', VAOPS, 'Cancelled Campaign'],
  ['vam-rca', 'va-managers', 'clients', 'RCA 2026', VAOPS, 'RCA 2026'],
  ['sales-hunter', 'sales', 'clients', 'Sales hunter forecast', MASTER, '5. Sales Hunter Forecast'],
  ['sales-forecast', 'sales', 'clients', 'Forecast', MASTER, '4. Forecast'],
  ['sales-dean', 'sales', 'clients', 'Bookings (Dean)', MASTER, 'Book from Dean'],
  ['mkt-va-tab', 'marketing', 'leads', 'Daily leads (VA tab)', MKT, 'VA Tab'],
  ['mkt-monthly', 'marketing', 'leads', 'Monthly summary', MKT, 'Monthly Summary'],
  ['mkt-plan', 'marketing', 'leads', 'Top 5 working plan', MKT, 'Top 5 Working Plan'],
  ['mkt-ads-cpl', 'marketing', 'leads', 'Ad spend & CPL', MASTER, 'Ads Spent CPL'],
]
PART_BYTES = 2_500_000

def slug(s):
    s = re.sub(r'[^a-z0-9]+', '-', str(s).lower()).strip('-')
    return s[:50] or 'tab'

# ---------- display formatting (approximates what Google Sheets shows) ----------
def first_section(fmt, value):
    parts = fmt.split(';')
    if len(parts) > 1 and isinstance(value, (int, float)) and value < 0:
        return parts[1], True
    if len(parts) > 2 and isinstance(value, (int, float)) and value == 0:
        return parts[2], False
    return parts[0], False

DATE_TOKENS = [('yyyy', '%Y'), ('yy', '%y'), ('mmmm', '%B'), ('mmm', '%b'), ('dddd', '%A'), ('ddd', '%a'), ('dd', '%d'), ('hh', '%H'), ('ss', '%S')]

def fmt_date(v, fmt):
    f = re.sub(r'\[[^\]]*\]', '', fmt or '')
    f = f.replace('\\', '').replace('"', '')
    if not re.search(r'[dmyhs]', f, re.I) or fmt in ('General', '@'):
        f = 'dd/mm/yyyy' if isinstance(v, (datetime.datetime, datetime.date)) else f
    low = f.lower()
    ampm = 'am/pm' in low
    low = low.replace('am/pm', '%p')
    out, i = '', 0
    while i < len(low):
        matched = False
        for tok, rep in DATE_TOKENS:
            if low.startswith(tok, i):
                out += rep; i += len(tok); matched = True; break
        if matched:
            continue
        ch = low[i]
        if ch == 'm':
            n = 2 if low.startswith('mm', i) else 1
            prev = out.rstrip()[-2:] if out else ''
            # m after an hour = minutes
            out += ('%M' if prev in ('%H', '%I') or ':' in out[-1:] else ('%m' if n == 2 else '%-m'))
            i += n
        elif ch == 'd':
            out += '%-d'; i += 1
        elif ch == 'h':
            out += '%-I' if ampm else '%-H'; i += 1
        elif ch == 's':
            out += '%S'; i += 1
        else:
            out += ch; i += 1
    try:
        if isinstance(v, datetime.time):
            v = datetime.datetime.combine(datetime.date(2000, 1, 1), v)
        if isinstance(v, datetime.date) and not isinstance(v, datetime.datetime):
            v = datetime.datetime.combine(v, datetime.time())
        return v.strftime(out)
    except Exception:
        return str(v)

def fmt_number(v, fmt):
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if fmt in (None, '', 'General', '@'):
        if isinstance(v, float):
            if math.isfinite(v) and v == int(v) and abs(v) < 1e15:
                return str(int(v))
            return ('%.10g' % v)
        return str(v)
    sec, neg_section = first_section(fmt, v)
    if re.search(r'[dmy]', re.sub(r'\[[^\]]*\]|"[^"]*"', '', sec), re.I) and not re.search(r'[0#]', sec):
        return str(v)
    prefix = ''.join(m.group(1) for m in re.finditer(r'\[\$([^\-\]]+)[^\]]*\]', sec))
    s = re.sub(r'\[[^\]]*\]', '', sec)
    lit = re.findall(r'"([^"]*)"', s)
    core = re.sub(r'"[^"]*"', '\x00', s)
    pct = '%' in core
    m = re.search(r'[#0][#0,]*(\.[0#]+)?', core)
    if not m:
        return str(v)
    dec = len(m.group(1)) - 1 if m.group(1) else 0
    comma = ',' in m.group(0)
    x = abs(v) if (neg_section or v < 0) else v
    if pct:
        x = x * 100
    body = f'{x:,.{dec}f}' if comma else f'{x:.{dec}f}'
    before = core[:m.start()].replace('\x00', lit[0] if lit else '', 1)
    after = core[m.end():]
    if '\x00' in after:
        after = after.replace('\x00', lit[-1] if lit else '', 1)
    before = re.sub(r'[_*\\].?', '', before).replace('(', '(').strip() if before else ''
    after = re.sub(r'[_*\\].?', '', after)
    sign = '-' if (v < 0 and not neg_section) else ''
    text = f'{sign}{prefix}{before}{body}{after}'
    if neg_section and not ('(' in sec):
        text = '-' + text if not text.startswith('-') else text
    return text.strip()

def display(cell):
    v = cell.value
    if v is None:
        return ''
    if isinstance(v, str):
        return v
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return fmt_date(v, cell.number_format)
    if isinstance(v, (int, float)):
        try:
            return fmt_number(v, cell.number_format)
        except Exception:
            return str(v)
    return str(v)

def color(c):
    try:
        if c is None or c.type != 'rgb' or not isinstance(c.rgb, str) or len(c.rgb) < 6:
            return None
        return '#' + c.rgb[-6:].lower()
    except Exception:
        return None

def style_of(cell):
    st = {}
    f = cell.fill
    if f is not None and f.fill_type == 'solid':
        bg = color(f.fgColor)
        if bg and bg != '#ffffff':
            st['bg'] = bg
    fo = cell.font
    if fo is not None:
        fg = color(fo.color)
        if fg and fg != '#000000':
            st['fg'] = fg
        if fo.bold:
            st['b'] = 1
        if fo.italic:
            st['i'] = 1
    al = cell.alignment
    if al is not None and al.horizontal in ('center', 'centerContinuous'):
        st['a'] = 'c'
    elif al is not None and al.horizontal == 'right':
        st['a'] = 'r'
    return st

def hidden_cols(ws):
    hid, widths = set(), {}
    for key, dim in ws.column_dimensions.items():
        lo, hi = dim.min or 0, dim.max or 0
        if not lo:
            continue
        for c in range(lo, min(hi, 2000) + 1):
            if dim.hidden:
                hid.add(c)
            if dim.width:
                widths[c] = dim.width
    return hid, widths

def convert(t, position, wb, file_id):
    key, dept, area, label, _, tab = t[:6]
    opt = t[6] if len(t) > 6 else {}
    ws = wb[opt.get('xlsx', tab)]  # Excel export drops '/' from tab names
    max_r, max_c = ws.max_row or 1, ws.max_column or 1
    hcols, widths = hidden_cols(ws)
    keep_cols = [c for c in range(1, max_c + 1) if c not in hcols]
    head = opt.get('headerRows', 1)
    if opt.get('lastRows') and max_r > opt['lastRows'] + head:
        row_ids = list(range(1, head + 1)) + list(range(max_r - opt['lastRows'] + 1, max_r + 1))
    else:
        row_ids = list(range(1, max_r + 1))
    notes_seen = 0
    rows_in = []
    for r in row_ids:
        hidden = bool(ws.row_dimensions[r].hidden) if r in ws.row_dimensions else False
        cells = []
        for c in keep_cols:
            cell = ws.cell(row=r, column=c)
            val = display(cell)
            note = cell.comment.text.strip() if cell.comment is not None and cell.comment.text else ''
            if note:
                notes_seen += 1
            cells.append((val, note, style_of(cell)))
        rows_in.append((r, hidden, cells))
    # trim empty rows at the end, and empty columns at the right
    has_text = lambda cells: any(v or n for v, n, _ in cells)
    while rows_in and not has_text(rows_in[-1][2]):
        rows_in.pop()
    width = 0
    for _, _, cells in rows_in:
        for i, (v, n, _) in enumerate(cells):
            if v or n:
                width = max(width, i + 1)
    width = max(width, 1)
    cols = keep_cols[:width]
    parts, cur = [], None
    def new_part():
        nonlocal cur
        cur = {'styles': [], 'index': {}, 'rows': [], 'bytes': 0}
        parts.append(cur)
    new_part()
    row_pos = {}
    for r, hidden, cells in rows_in:
        cells = cells[:width]
        end = len(cells)
        while end > 0 and not cells[end - 1][0] and not cells[end - 1][1] and not cells[end - 1][2]:
            end -= 1
        cells = cells[:end]
        est = len(json.dumps(cells))
        if cur['rows'] and cur['bytes'] + est > PART_BYTES:
            new_part()
        row = [-r if hidden else r]
        for v, n, st in cells:
            idx = -1
            if st:
                sk = json.dumps(st, sort_keys=True)
                if sk not in cur['index']:
                    cur['index'][sk] = len(cur['styles']); cur['styles'].append(st)
                idx = cur['index'][sk]
            if not v and not n and idx < 0:
                row.append(None)
            elif idx < 0 and not n:
                row.append(v)
            elif not n:
                row.append([v, idx])
            else:
                row.append([v, idx, n[:1000]])
        row_pos[r] = len(row_pos)
        cur['rows'].append(row)
        cur['bytes'] += est
    col_pos = {c: i for i, c in enumerate(cols)}
    merges = []
    for mr in ws.merged_cells.ranges:
        r0, c0 = row_pos.get(mr.min_row), col_pos.get(mr.min_col)
        if r0 is None or c0 is None:
            continue
        rs = sum(1 for rr in range(mr.min_row, mr.max_row + 1) if rr in row_pos)
        cs = sum(1 for cc in range(mr.min_col, mr.max_col + 1) if cc in col_pos)
        if rs * cs > 1:
            merges.append([r0, c0, rs, cs])
    fr, fc = 0, 0
    if ws.freeze_panes:
        cell = ws[ws.freeze_panes]
        fr, fc = cell.row - 1, cell.column - 1
    frozen_rows = sum(1 for r in row_pos if r <= fr) or min(1, len(row_pos))
    frozen_cols = sum(1 for c in cols if c <= fc)
    meta = {
        'key': key, 'department': dept, 'area': area, 'label': label, 'position': position,
        'source_name': f"{NAMES[file_id]} › {tab.strip()}",
        'source_url': (f'https://docs.google.com/spreadsheets/d/{file_id}/edit' if file_id != ROEL else f'https://drive.google.com/file/d/{file_id}/view'),
        'frozen_rows': frozen_rows, 'frozen_cols': frozen_cols,
        'col_widths': [round(widths.get(c, 12.6) * 7 + 5) for c in cols],
        'merges': merges, 'row_count': len(row_pos), 'col_count': len(cols),
    }
    out = {'meta': meta, 'parts': [{'styles': p['styles'], 'rows': p['rows']} for p in parts]}
    json.dump(out, open(os.path.join(OUT, key + '.json'), 'w'))
    return len(row_pos), len(cols), len(parts), notes_seen

books = {}
def book(fid):
    if fid not in books:
        books[fid] = openpyxl.load_workbook(os.path.join(SRC_DIR, fid + '.xlsx'), data_only=True)
    return books[fid]

tabs = []
for t in TABS:
    if t == 'ROEL':
        wb = book(ROEL)
        for ws in wb.worksheets:
            if ws.sheet_state != 'visible':
                continue
            tabs.append(['rec-' + slug(ws.title), 'recruitment', 'candidates', ws.title, ROEL, ws.title])
    else:
        tabs.append(t)
ok = 0
skipped = []
for i, t in enumerate(tabs):
    try:
        n, w, p, notes = convert(t, i, book(t[4]), t[4])
        if n == 0:
            os.remove(os.path.join(OUT, t[0] + '.json'))
            skipped.append(t[0])
            print(f'SKIP {t[0]:<34} empty tab')
            continue
        ok += 1
        print(f'OK   {t[0]:<34} {n:>5} rows x {w:>3} cols, {p} part(s), {notes} notes')
    except Exception as e:
        print(f'FAIL {t[0]:<34} {type(e).__name__}: {str(e)[:120]}')
json.dump([t[0] for t in tabs if t[0] not in skipped and os.path.exists(os.path.join(OUT, t[0] + '.json'))], open(os.path.join(OUT, '_keys.json'), 'w'))
print(f'{ok} of {len(tabs)} tabs converted')
