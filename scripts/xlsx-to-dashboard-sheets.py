"""Turn every visible tab of an exported Google Sheet (.xlsx) into an editable dashboard Sheet (JSON),
for teams that move off the Google Sheet. One-way: nothing is ever written back to Google.

usage: python3 scripts/xlsx-to-dashboard-sheets.py <file.xlsx> <out dir> "<source name>"
then:  npx tsx scripts/load-dashboard-sheets.mts <out dir> <department key>

Per tab: the header is the row (in the first 6) with the most filled cells; a blank header cell takes the
text above it (e.g. "REMARKS/NOTES"), else "Column N". Values are kept as the sheet shows them (€, %, dates).
Hidden columns are left out; hidden rows are kept (they are still data). Cell notes go into a
"Sheet notes" column at the end ("Column: note | ..."). Empty tabs are skipped.
"""
import importlib.util, json, os, sys

here = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("mirror_tools", os.path.join(here, "xlsx-to-sheet-mirror.py"))
tools = importlib.util.module_from_spec(spec)
spec.loader.exec_module(tools)  # only defines helpers (main() is guarded)
import openpyxl

src, out, source_name = sys.argv[1], sys.argv[2], sys.argv[3]
os.makedirs(out, exist_ok=True)
wb = openpyxl.load_workbook(src, data_only=True)
index = []
for ws in wb.worksheets:
    if ws.sheet_state != "visible":
        continue
    max_r, max_c = ws.max_row or 1, ws.max_column or 1
    hidden, _ = tools.hidden_cols(ws)
    cols = [c for c in range(1, max_c + 1) if c not in hidden]
    grid = []  # (values, notes) per row
    for r in range(1, max_r + 1):
        vals, notes = [], []
        for c in cols:
            cell = ws.cell(row=r, column=c)
            vals.append(tools.display(cell).strip())
            notes.append(cell.comment.text.strip() if cell.comment is not None and cell.comment.text else "")
        grid.append((vals, notes))
    while grid and not any(v for v in grid[-1][0]) and not any(n for n in grid[-1][1]):
        grid.pop()
    if not grid:
        print(f"SKIP {ws.title!r}: empty")
        continue
    # keep only columns that have something in them
    used = [i for i in range(len(cols)) if any(row[0][i] or row[1][i] for row in grid)]
    if not used:
        print(f"SKIP {ws.title!r}: empty")
        continue
    h = max(range(min(6, len(grid))), key=lambda i: (sum(1 for j in used if grid[i][0][j]), -i))
    names, seen = [], {}
    for j in used:
        name = grid[h][0][j] or next((grid[k][0][j] for k in range(h - 1, -1, -1) if grid[k][0][j]), "") or f"Column {tools.get_column_letter(cols[j])}"
        name = " ".join(name.split())[:72] or f"Column {j + 1}"
        if name.lower() in seen:
            seen[name.lower()] += 1
            name = f"{name} ({seen[name.lower()]})"
        else:
            seen[name.lower()] = 1
        names.append(name)
    rows = []
    for vals, notes in grid[h + 1:]:
        cells = {str(k): vals[j] for k, j in enumerate(used) if vals[j]}
        note = " | ".join(f"{names[k]}: {notes[j]}" for k, j in enumerate(used) if notes[j])
        if note:
            cells["notes"] = note[:4000]
        rows.append(cells)  # empty rows are kept: they separate blocks (e.g. pay-runs) like in the sheet
    has_notes = any("notes" in r for r in rows)
    widths = [min(max(len(n) * 8 + 24, 90), 320) for n in names]
    file = f"{len(index):02d}.json"
    json.dump({"name": ws.title.strip()[:120], "source": f"{source_name} › {ws.title.strip()}", "columns": names,
               "widths": widths, "has_notes": has_notes, "rows": rows}, open(os.path.join(out, file), "w"))
    index.append(file)
    print(f"OK   {ws.title.strip()!r:<40} {len(rows):>5} rows x {len(names):>3} cols{' + notes' if has_notes else ''}")
json.dump(index, open(os.path.join(out, "_index.json"), "w"))
print(f"{len(index)} tabs")
